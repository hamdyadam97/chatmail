"""Excel export of an aggregate report (the same dict build_report returns).

Only counts and averages — the export cannot carry more than the on-screen
report, and access/auditing are those of the report view. Every text cell is
written as an explicit string so a unit name such as "=HYPERLINK(...)" can
never become a formula in the reader's Excel (CSV/formula injection).
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.approvals.models import ApprovalRequest

HEADER_FILL = PatternFill("solid", fgColor="1E3A8A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=12)

APPROVAL_STATUS = dict(ApprovalRequest.STATUSES)


def _put(ws, row, col, value, *, font=None, fill=None):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # never a formula
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    cell.alignment = Alignment(horizontal="right" if isinstance(value, str) else "center")
    return cell


def _blank(value):
    return "—" if value is None else value


def _pct(ratio):
    return None if ratio is None else round(ratio * 100, 1)


def _sheet(wb, title, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    return ws


def _header(ws, row, labels):
    for col, label in enumerate(labels, start=1):
        _put(ws, row, col, label, font=HEADER_FONT, fill=HEADER_FILL)


def build_workbook(report: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    c, a, n, u = report["correspondence"], report["approvals"], report["announcements"], report["usage"]

    ws = _sheet(wb, "ملخص", [38, 22])
    rows = [
        ("النطاق", report["scope"]["name"]),
        ("الفترة", f'{report["period"]["from"]} — {report["period"]["to"]}'),
        ("تاريخ الإنشاء", report["generated_at"][:16].replace("T", " ")),
        None,
        ("المراسلات", None),
        ("المرسلة", c["sent"]),
        ("خيوط جديدة", c["new_threads"]),
        ("ردود", c["replies"]),
        ("إحالات", c["forwards"]),
        ("مغلقة", c["closed"]),
        ("متأخرة (مفتوحة بعد موعدها)", c["overdue_open"]),
        ("لها موعد رد", c["with_due_date"]),
        ("متوسط زمن أول رد (ساعة)", _blank(c["avg_first_reply_hours"])),
        ("وسيط زمن أول رد (ساعة)", _blank(c["median_first_reply_hours"])),
        None,
        ("الاعتمادات", None),
        ("طلبات الاعتماد", a["requests"]),
        *[(f"— {APPROVAL_STATUS.get(k, k)}", v) for k, v in a["by_status"].items()],
        ("متوسط زمن القرار (ساعة)", _blank(a["avg_decision_hours"])),
        None,
        ("التعاميم", None),
        ("المنشورة", n["published"]),
        ("تتطلب إقرار اطلاع", n["requiring_ack"]),
        ("نسبة الإقرار %", _blank(_pct(n["ack_rate"]))),
        ("إقرارات متأخرة", n["overdue_acks"]),
        None,
        ("الاستخدام", None),
        ("الموظفون", u["employees"]),
        ("نشطون آخر 30 يومًا", u["active_last_30_days"]),
        ("أجهزة موثّقة نشطة", u["active_devices"]),
        ("رسائل المحادثات", u["chat_messages"]),
    ]
    r = 1
    for item in rows:
        if item is None:
            r += 1
            continue
        label, value = item
        if value is None:
            _put(ws, r, 1, label, font=SECTION_FONT)
        else:
            _put(ws, r, 1, label)
            _put(ws, r, 2, value)
        r += 1

    ws = _sheet(wb, "شهري", [14, 16])
    _header(ws, 1, ["الشهر", "المراسلات المرسلة"])
    for i, row in enumerate(report["monthly"], start=2):
        _put(ws, i, 1, row["month"])
        _put(ws, i, 2, row["sent"])

    b = report["breakdown"]
    ws = _sheet(wb, "حسب الوحدة", [30, 12, 12, 12, 12, 14, 22])
    _header(ws, 1, ["الوحدة", "الرمز", "الموظفون", "المرسلة", "المغلقة", "المتأخرة", "متوسط أول رد (ساعة)"])
    r = 2
    for row in b["rows"]:
        for col, value in enumerate(
            [row["unit"]["name"], row["unit"]["code"], row["employees"], row["sent"], row["closed"],
             row["overdue_open"], _blank(row["avg_first_reply_hours"])],
            start=1,
        ):
            _put(ws, r, col, value)
        r += 1
    merged = b["merged_small_units"]
    if merged:
        _put(ws, r, 1, f'وحدات صغيرة مدمجة ({merged["units"]}) — أقل من {b["min_group_size"]} موظفين لكل منها')
        _put(ws, r, 3, merged["employees"])
        _put(ws, r, 4, merged["sent"])
        _put(ws, r, 6, merged["overdue_open"])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
