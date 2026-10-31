"""Bulk-create employee accounts from the HR spreadsheet.

    manage.py import_employees FILE.xlsx                 # dry run: report only
    manage.py import_employees FILE.xlsx --apply --out /root/passwords.xlsx

Columns (first sheet, by header): الاسم، البريد الإلكتروني، الفرع، المعهد،
ملاحظة. The username is the part of the e-mail before "@"; the e-mail itself
is not stored (user decision). Each account gets a one-time temporary password
(must be changed at first login), a primary membership in its branch and the
employee role there (grant_default_role). Rows are skipped, with the reason in
the report, when: the username already exists, it repeats in the file, it is
marked as a test account, or it is a branch-less "possible duplicate". Rows
without a branch are created without a membership (no role yet) — assign
their branch from the employees page.

The --out file holds the temporary passwords: it is written 0600 and must be
handed out and deleted. Nothing is written in a dry run.
"""
import os
import re
from collections import Counter

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

from apps.accounts.models import User
from apps.accounts.password_policy import generate_temp_password
from apps.audit.services import log
from apps.authorization.services import grant_default_role
from apps.org.models import Branch, EmployeeMembership, Institute
from apps.textsearch import normalize

COLUMNS = {"name": "الاسم", "email": "البريد الإلكتروني", "branch": "الفرع", "institute": "المعهد", "note": "ملاحظة"}
USERNAME_RE = re.compile(r"^[A-Za-z0-9._-]{1,64}$")


def _key(text: str) -> str:
    """Branch/institute names compared loosely: Arabic normalization, no
    dashes/spaces, without the words فرع/معهد."""
    text = normalize(text or "")
    text = re.sub(r"\b(فرع|معهد)\b", " ", text)
    return re.sub(r"[\s\-—–_]+", "", text)


def _find_branch(branches: dict, name: str):
    """Exact (normalized) name first, then without the "— رجال/نساء" suffix,
    then a unique system branch whose name is contained in it."""
    key = _key(name)
    if key in branches:
        return branches[key]
    bare = _key(re.sub(r"[—–-]\s*(رجال|نساء)\s*$", "", name or ""))
    if bare in branches:
        return branches[bare]
    hits = [b for k, b in branches.items() if k and k in key]
    return hits[0] if len(hits) == 1 else None


class Command(BaseCommand):
    help = "Import employees from the HR spreadsheet (dry run unless --apply)."

    def add_arguments(self, parser):
        parser.add_argument("file")
        parser.add_argument("--apply", action="store_true", help="really create the accounts")
        parser.add_argument("--out", default="", help="xlsx report with the temporary passwords")
        parser.add_argument("--by", default="it.admin", help="administrator recorded as the creator")
        parser.add_argument("--include-test", action="store_true", help="also import rows marked as test accounts")

    def handle(self, *args, **opts):
        import openpyxl

        admin = User.objects.filter(username=opts["by"]).first()
        if admin is None:
            raise CommandError(f"no user {opts['by']!r} (use --by)")
        if opts["apply"] and not opts["out"]:
            raise CommandError("--apply needs --out (the temporary passwords go there)")

        sheet = openpyxl.load_workbook(opts["file"], read_only=True).worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        header = [str(h or "").strip() for h in rows[0]]
        try:
            idx = {k: header.index(v) for k, v in COLUMNS.items()}
        except ValueError as exc:
            raise CommandError(f"missing column: {exc}")

        branches = {_key(b.name): b for b in Branch.objects.select_related("institute")}
        institutes = {_key(i.name): i for i in Institute.objects.all()}
        records = []
        for r in rows[1:]:
            get = lambda k: str(r[idx[k]] or "").strip() if idx[k] < len(r) else ""  # noqa: E731
            email = get("email")
            if not email:
                continue
            records.append({
                "name": get("name"), "username": email.split("@", 1)[0].strip().lower(),
                "branch": get("branch"), "institute": get("institute"), "note": get("note"),
            })
        seen = Counter(rec["username"] for rec in records)
        existing = set(User.objects.filter(username__in=list(seen)).values_list("username", flat=True))

        plan, missing_branches, done = [], Counter(), set()
        for rec in records:
            u, status, branch = rec["username"], "", None
            if not USERNAME_RE.match(u):
                status = "اسم مستخدم غير صالح"
            elif u in existing:
                status = "موجود مسبقًا"
            elif u in done:
                status = "مكرر في الملف"
            elif "تجربة" in rec["note"] and not opts["include_test"]:
                status = "حساب تجربة — لم يُستورد"
            elif "مكرر محتمل" in rec["note"] and not rec["branch"]:
                status = "مكرر محتمل بلا فرع — لم يُستورد"
            elif rec["branch"]:
                branch = _find_branch(branches, rec["branch"])
                if branch is None:
                    missing_branches[rec["branch"]] += 1
                    status = "الفرع غير موجود في النظام"
                elif rec["institute"] and institutes.get(_key(rec["institute"])) not in (None, branch.institute):
                    status = "الفرع لا يتبع المعهد المكتوب"
            if not status:
                status = "سيُنشأ" if branch else "سيُنشأ بلا فرع"
                done.add(u)
            plan.append({**rec, "status": status, "branch_obj": branch})

        for status, n in Counter(p["status"] for p in plan).most_common():
            self.stdout.write(f"{n:4d}  {status}")
        for name, n in missing_branches.most_common():
            self.stdout.write(f"      فرع غير موجود: {name} ({n})")
        if not opts["apply"]:
            self.stdout.write("dry run — nothing written (add --apply --out FILE)")
            return

        created = 0
        with transaction.atomic():
            for p in plan:
                if not p["status"].startswith("سيُنشأ"):
                    continue
                password = generate_temp_password(p["username"])
                user = User(username=p["username"], first_name=p["name"], full_name=p["name"] or p["username"],
                            must_change_password=True)
                user.set_password(password)
                user.save()
                if p["branch_obj"] is not None:
                    EmployeeMembership.objects.create(
                        user=user, institute=p["branch_obj"].institute, branch=p["branch_obj"],
                        is_primary=True, started_on=timezone.localdate(),
                    )
                    grant_default_role(user, p["branch_obj"].institute, granted_by=admin)
                log(actor=admin, action="account_created", target_type="user", target_id=user.id,
                    metadata={"via": "import_employees", "has_branch": p["branch_obj"] is not None})
                p["password"] = password
                p["status"] = "أُنشئ" if p["branch_obj"] else "أُنشئ بلا فرع"
                created += 1
        self._write_report(opts["out"], plan)
        self.stdout.write(self.style.SUCCESS(f"created {created} account(s); report: {opts['out']} (delete it after use)"))

    @staticmethod
    def _write_report(path, plan):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "الحسابات"
        ws.sheet_view.rightToLeft = True
        ws.append(["الاسم", "اسم المستخدم", "كلمة المرور المؤقتة", "الفرع", "الحالة"])
        for p in plan:
            # every cell as text (no formula injection from names)
            ws.append([f"{v}" for v in (p["name"], p["username"], p.get("password", ""), p["branch"], p["status"])])
            for cell in ws[ws.max_row]:
                cell.data_type = "s"
        for col, width in zip("ABCDE", (30, 34, 20, 36, 30)):
            ws.column_dimensions[col].width = width
        fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
        with os.fdopen(fd, "wb") as fh:
            wb.save(fh)
