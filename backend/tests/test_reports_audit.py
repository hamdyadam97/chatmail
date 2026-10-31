"""Management reports (aggregates only, scoped, small units merged, audited)
and the audit-log API (central audit.view; own security events for everyone)."""
import json
from datetime import timedelta

import pytest
from django.utils import timezone

from apps.audit.models import AuditLog

from .conftest import DEFAULT_PASSWORD, assign_role
from .correspondence_utils import (  # noqa: F401  (fixtures)
    CORR_URL,
    _blob_root,
    make_employee,
    office,
    prepare_and_send,
    reply,
)
from .messaging_utils import auth

pytestmark = pytest.mark.django_db

REPORT_URL = "/api/v1/reports/"
AUDIT_URL = "/api/v1/audit-logs/"
CANARY = "CANARY-report-subject-55aa"


@pytest.fixture
def activity(api_client, user_factory, office, branch_a, branch_a2):
    """office (4 employees in branch_a) exchanges correspondence; a manager
    of branch_a; one lone employee in branch_a2 (a 'small unit')."""
    o = office
    yesterday = (timezone.localdate() - timedelta(days=1)).isoformat()
    overdue = prepare_and_send(api_client, o["h_s"], [o["r1"]], subject=CANARY, due_date=yesterday)
    answered = prepare_and_send(api_client, o["h_s"], [o["r1"]], [o["r2"]], body=CANARY)
    assert reply(api_client, o["h_r1"], answered["id"], body=CANARY).status_code == 201
    assert api_client.post(f"{CORR_URL}{answered['id']}/close/", **o["h_s"]).status_code == 200
    manager = make_employee(user_factory, "rep_mgr", branch_a)
    assign_role(manager, "branch_manager", institute=branch_a.institute, branch=branch_a)
    lone = make_employee(user_factory, "rep_lone", branch_a2)
    o.update(manager=manager, h_mgr=auth(api_client, "rep_mgr"), lone=lone, overdue=overdue, answered=answered)
    return o


def _report(client, headers, unit_type, unit_id=None):
    url = f"{REPORT_URL}?unit_type={unit_type}" + (f"&unit_id={unit_id}" if unit_id else "")
    return client.get(url, **headers)


# --- reports: access ------------------------------------------------------------------

def test_report_access_is_scoped(api_client, user_factory, activity, branch_a, branch_a2, institute_a):
    a = activity
    assert _report(api_client, a["h_s"], "branch", branch_a.id).status_code == 403  # plain employee
    assert _report(api_client, a["h_mgr"], "branch", branch_a.id).status_code == 200
    assert _report(api_client, a["h_mgr"], "branch", branch_a2.id).status_code == 403
    assert _report(api_client, a["h_mgr"], "institute", institute_a.id).status_code == 403
    assert _report(api_client, a["h_mgr"], "all").status_code == 403
    exe = make_employee(user_factory, "rep_exec", branch_a)
    assign_role(exe, "executive")  # central reports.view_metadata
    assert _report(api_client, auth(api_client, "rep_exec"), "all").status_code == 200
    # sys_admin manages the system but holds no reports permission
    admin = user_factory("rep_admin", membership_unit=branch_a)
    assign_role(admin, "sys_admin")
    assert _report(api_client, auth(api_client, "rep_admin"), "all").status_code == 403


def test_units_selector_lists_only_permitted_units(api_client, activity, branch_a):
    units = api_client.get(f"{REPORT_URL}units/", **activity["h_mgr"]).json()
    assert {"type": "branch", "id": str(branch_a.id)} == {k: units[0][k] for k in ("type", "id")}
    assert all(u["type"] in ("branch", "department") for u in units)


# --- reports: content ---------------------------------------------------------------------

def test_metrics_count_threads_replies_overdue_and_closures(api_client, activity, branch_a):
    rep = _report(api_client, activity["h_mgr"], "branch", branch_a.id).json()
    c = rep["correspondence"]
    assert c["sent"] == 3 and c["new_threads"] == 2 and c["replies"] == 1
    assert c["closed"] == 1
    assert c["overdue_open"] == 1
    assert c["replied_threads"] == 1 and c["avg_first_reply_hours"] is not None
    assert rep["usage"]["employees"] == 5  # office (4) + manager
    assert sum(m["sent"] for m in rep["monthly"]) == 3


def test_report_carries_no_identities_or_content(api_client, activity, branch_a):
    body = json.dumps(_report(api_client, activity["h_mgr"], "branch", branch_a.id).json(), ensure_ascii=False)
    assert CANARY not in body
    for key in ("s", "r1", "r2", "o", "manager"):
        user = activity[key]
        assert user.username not in body and user.full_name not in body
        assert str(user.id) not in body
    assert activity["answered"]["reference_number"] not in body


def test_small_units_are_merged(api_client, user_factory, activity, institute_a, branch_a):
    head = make_employee(user_factory, "rep_inst", branch_a)
    assign_role(head, "institute_manager", institute=institute_a)
    rep = _report(api_client, auth(api_client, "rep_inst"), "institute", institute_a.id).json()
    names = [r["unit"]["code"] for r in rep["breakdown"]["rows"]]
    assert "RYD" in names and "JED" not in names  # JED has a single employee
    assert rep["breakdown"]["merged_small_units"]["units"] >= 1


def test_report_access_is_audited(api_client, activity, branch_a):
    _report(api_client, activity["h_mgr"], "branch", branch_a.id)
    row = AuditLog.objects.filter(action="report_viewed", actor=activity["manager"]).get()
    assert row.target_id == branch_a.id


def test_bad_period_is_rejected(api_client, activity, branch_a):
    url = f"{REPORT_URL}?unit_type=branch&unit_id={branch_a.id}&from=2020-01-01&to=2026-01-01"
    assert api_client.get(url, **activity["h_mgr"]).status_code == 400


def test_health_for_report_viewers_only(api_client, activity):
    assert api_client.get(f"{REPORT_URL}health/", **activity["h_s"]).status_code == 403
    health = api_client.get(f"{REPORT_URL}health/", **activity["h_mgr"]).json()
    assert set(health["checks"]) == {"database", "cache", "realtime", "blob_storage", "content_key"}
    assert health["checks"]["database"]["ok"] is True


# --- audit log API ------------------------------------------------------------------------

def test_full_audit_log_needs_central_audit_view(api_client, user_factory, activity, branch_a):
    auditor = user_factory("aud_sec", membership_unit=branch_a)
    assign_role(auditor, "security_auditor")
    h = auth(api_client, "aud_sec")
    page = api_client.get(f"{AUDIT_URL}?action=correspondence_sent", **h).json()
    assert page["count"] == 3 and {r["action"] for r in page["results"]} == {"correspondence_sent"}
    assert AuditLog.objects.filter(action="audit_log_viewed", actor=auditor).exists()
    assert "correspondence_sent" in api_client.get(f"{AUDIT_URL}actions/", **h).json()
    for headers in (activity["h_mgr"], activity["h_s"]):
        assert api_client.get(AUDIT_URL, **headers).status_code == 403
    admin = user_factory("aud_admin", membership_unit=branch_a)
    assign_role(admin, "sys_admin")
    assert api_client.get(AUDIT_URL, **auth(api_client, "aud_admin")).status_code == 403


def test_my_security_events(api_client, user_factory, branch_a):
    me = user_factory("aud_me", membership_unit=branch_a)
    user_factory("aud_other")
    auth(api_client, "aud_other")
    h = auth(api_client, "aud_me")
    api_client.post("/api/v1/auth/login/", {"username": "aud_me", "password": "wrong"}, format="json")
    rows = api_client.get(f"{AUDIT_URL}mine/", **h).json()
    actions = [r["action"] for r in rows]
    assert "login" in actions
    assert all(r["actor"] is None or r["actor"]["id"] == str(me.id) for r in rows)
    assert not any(r["actor"] and r["actor"]["username"] == "aud_other" for r in rows)
    assert DEFAULT_PASSWORD not in json.dumps(rows) and "wrong" not in json.dumps(rows)


# --- Excel export ---------------------------------------------------------------------

def _export(client, headers, unit_type, unit_id=None):
    url = f"{REPORT_URL}export/?unit_type={unit_type}" + (f"&unit_id={unit_id}" if unit_id else "")
    return client.get(url, **headers)


def test_export_matches_report_access_and_is_audited(api_client, activity, branch_a, branch_a2):
    a = activity
    assert _export(api_client, a["h_s"], "branch", branch_a.id).status_code == 403
    assert _export(api_client, a["h_mgr"], "branch", branch_a2.id).status_code == 403
    r = _export(api_client, a["h_mgr"], "branch", branch_a.id)
    assert r.status_code == 200
    assert r["Content-Type"].startswith("application/vnd.openxmlformats")
    assert AuditLog.objects.filter(action="report_exported", actor=a["manager"], target_id=branch_a.id).exists()


def test_export_content_is_the_aggregate_report_only(api_client, activity, branch_a):
    from io import BytesIO

    from openpyxl import load_workbook

    r = _export(api_client, activity["h_mgr"], "branch", branch_a.id)
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["ملخص", "شهري", "حسب الوحدة"]
    summary = {row[0]: row[1] for row in wb["ملخص"].iter_rows(values_only=True) if row[0]}
    report = _report(api_client, activity["h_mgr"], "branch", branch_a.id).json()
    assert summary["المرسلة"] == report["correspondence"]["sent"]
    assert summary["متأخرة (مفتوحة بعد موعدها)"] == report["correspondence"]["overdue_open"]
    flat = repr([list(ws.iter_rows(values_only=True)) for ws in wb.worksheets])
    assert CANARY not in flat
    for name in ("c_sender", "c_rcpt1", "rep_mgr"):
        assert name not in flat  # no identities, counts only


def test_export_never_writes_formulas(api_client, user_factory, activity, institute_a, branch_a):
    """A unit name crafted as a formula stays inert text in Excel."""
    from io import BytesIO

    from openpyxl import load_workbook

    branch_a.name = '=HYPERLINK("http://evil.example","x")'
    branch_a.save(update_fields=["name"])
    head = make_employee(user_factory, "rep_inst2", branch_a)
    assign_role(head, "institute_manager", institute=institute_a)
    r = _export(api_client, auth(api_client, "rep_inst2"), "institute", institute_a.id)
    ws = load_workbook(BytesIO(r.content))["حسب الوحدة"]
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert cells and all(c.data_type == "s" for c in cells)
