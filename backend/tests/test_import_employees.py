"""import_employees: dry run writes nothing; apply creates accounts in their
branch with the employee role, a one-time password, and skips what it should."""
import openpyxl
import pytest
from django.core.management import call_command

from apps.accounts.models import User

pytestmark = pytest.mark.django_db
HEADER = ("الاسم", "البريد الإلكتروني", "الدور", "الفرع", "المعهد", "القسم", "الحالة", "آخر دخول", "ملاحظة")


@pytest.fixture
def sheet(tmp_path, branch_a, user_factory):
    from .conftest import assign_role

    assign_role(user_factory("it.admin"), "sys_admin")
    user_factory("taken")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADER)
    b, i = f"فرع {branch_a.name} — رجال", branch_a.institute.name
    ws.append(("أحمد كامل", "Ahmed.Kamal@ararhni.com", "موظف", b, i, None, "نشط", "", None))
    ws.append(("بلا فرع", "nobranch@ararhni.com", None, None, None, None, "نشط", "", "يحتاج فرع"))
    ws.append(("تجربة", "dev@root-jo.com", None, None, None, None, "نشط", "", "حساب تجربة/مبرمج؟"))
    ws.append(("مكرر", "ali_x@ararhni.com", None, None, None, None, "نشط", "", "مكرر محتمل مع ali.x@ararhni.com"))
    ws.append(("موجود", "taken@ararhni.com", "موظف", b, i, None, "نشط", "", None))
    ws.append(("فرع مجهول", "lost@ararhni.com", "موظف", "فرع غير موجود", i, None, "نشط", "", None))
    path = tmp_path / "staff.xlsx"
    wb.save(path)
    return path


def test_dry_run_then_apply(sheet, tmp_path, branch_a, capsys):
    call_command("import_employees", str(sheet))
    out = capsys.readouterr().out
    assert "dry run" in out and "فرع غير موجود" in out
    assert not User.objects.filter(username="ahmed.kamal").exists()

    report = tmp_path / "passwords.xlsx"
    call_command("import_employees", str(sheet), "--apply", "--out", str(report))
    ahmed = User.objects.get(username="ahmed.kamal")
    assert ahmed.full_name == "أحمد كامل" and ahmed.must_change_password and ahmed.optional_email is None
    assert ahmed.memberships.get().branch_id == branch_a.id
    assert ahmed.role_assignments.filter(role__code="employee", institute=branch_a.institute).exists()
    assert User.objects.filter(username="nobranch").exists()
    for skipped in ("dev", "ali_x", "lost"):
        assert not User.objects.filter(username=skipped).exists()

    rows = list(openpyxl.load_workbook(report).active.iter_rows(values_only=True))
    by_user = {r[1]: r for r in rows[1:]}
    password = by_user["ahmed.kamal"][2]
    assert ahmed.check_password(password) and by_user["ahmed.kamal"][4] == "أُنشئ"
    assert by_user["taken"][4] == "موجود مسبقًا" and not by_user["dev"][2]
